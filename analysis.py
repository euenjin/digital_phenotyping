from pathlib import Path
import shutil

import numpy as np
import pandas as pd
import statsmodels.api as sm


DATA_PATH = Path("df_cohort_cleaned.csv")
OUT_DIR = Path("regression_results")
MODEL_123_XLSX = OUT_DIR / "model_1_2_3.xlsx"

PHQ_CATEGORIES = [
    "Minimal",
    "Mild_or_greater",
]

PULSE_BINS = [-np.inf, 59, 69, 79, 89, np.inf]
PULSE_LABELS = ["<60 bpm", "60-69 bpm", "70-79 bpm", "80-89 bpm", ">=90 bpm"]
PULSE_REFERENCE = "<60 bpm"

SEX_LEVELS = ["Male", "Female"]
INCOME_LEVELS = ["Low", "Middle-low", "Middle-high", "High"]
EDUCATION_LEVELS = ["Primary school or below", "Middle school", "High school", "College or above"]
MARITAL_LEVELS = ["Yes", "No"]
WORKING_LEVELS = ["Yes", "No"]
SMOKING_LEVELS = ["Never", "Previous", "Current"]
ALCOHOL_LEVELS = ["Never", "Previous", "Current"]
PHYSICAL_ACTIVITY_LEVELS = ["<600 MET-min/week", ">=600 MET-min/week"]

REFERENCE_ROWS = {
    "pulse_rate_60_group": ("pulse_rate_60_group", PULSE_REFERENCE),
    "sex_group": ("sex", "Male"),
    "education_group": ("education", "College or above"),
    "income_group": ("income", "High"),
    "marital_status": ("marital_status", "Yes"),
    "working_status": ("working_status", "Yes"),
    "alcohol_group": ("alcohol", "Never"),
    "smoking_group": ("smoking", "Never"),
    "physical_activity_group": ("physical_activity", ">=600 MET-min/week"),
}

TERM_ORDER = {
    "const": (0, 0),
    "age_numeric": (1, 0),
    "sex": (2, 0),
    "education": (3, 0),
    "income": (4, 0),
    "marital_status": (5, 0),
    "working_status": (6, 0),
    "alcohol": (7, 0),
    "smoking": (8, 0),
    "physical_activity": (9, 0),
    "sleep_avg_weighted": (10, 0),
    "BMI_numeric": (11, 0),
    "pulse_rate_60_group": (12, 0),
}

LEVEL_ORDER = {
    "pulse_rate_60_group": PULSE_LABELS,
    "sex": SEX_LEVELS,
    "education": EDUCATION_LEVELS[::-1],
    "income": INCOME_LEVELS[::-1],
    "marital_status": MARITAL_LEVELS,
    "working_status": WORKING_LEVELS,
    "alcohol": ALCOHOL_LEVELS,
    "smoking": SMOKING_LEVELS,
    "physical_activity": PHYSICAL_ACTIVITY_LEVELS[::-1],
}


def prepare_data(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path)

    df["PHQ_group"] = pd.Categorical(
        df["PHQ_group"],
        categories=PHQ_CATEGORIES,
        ordered=True,
    )
    df["PHQ_code"] = df["PHQ_group"].cat.codes
    df["PHQ_binary"] = df["PHQ_group"].map({
        "Minimal": 0,
        "Mild_or_greater": 1,
    }).astype("Int64")

    df["pulse_rate_60_group"] = pd.cut(
        df["pulse_rate_60"],
        bins=PULSE_BINS,
        labels=PULSE_LABELS,
        right=True,
    )
    df["pulse_rate_60_per_10bpm"] = df["pulse_rate_60"] / 10

    df["sex_group"] = df["sex_numeric"].map({
        1: "Male",
        2: "Female",
    })
    df["income_group"] = df["household_income"].map({
        1: "Low",
        2: "Middle-low",
        3: "Middle-high",
        4: "High",
    })
    df["education_group"] = df["edu_numeric"].map({
        1: "Primary school or below",
        2: "Middle school",
        3: "High school",
        4: "College or above",
    })
    df["marital_status"] = df["marri_1_numeric"].map({
        1: "Yes",
        2: "No",
    })
    df["working_status"] = df["employment_status"].map({
        1: "Yes",
        2: "No",
    })
    df["smoking_group"] = df["smoking_history"].map({
        1: "Never",
        2: "Previous",
        3: "Current",
    })
    df["alcohol_group"] = df["BD1_11_numeric"].map({
        1: "Never",
        2: "Previous",
        3: "Current",
    })

    return df


def add_constant(x: pd.DataFrame) -> pd.DataFrame:
    return sm.add_constant(x.astype(float), has_constant="add")


def categorical_dummies(
    series: pd.Series,
    levels: list[str],
    reference: str,
    prefix: str,
) -> pd.DataFrame:
    columns = {}
    for level in levels:
        if level == reference:
            continue
        columns[f"{prefix}[{level}]"] = (series == level).astype(float)
    return pd.DataFrame(columns, index=series.index)


def exposure_design(complete: pd.DataFrame, exposure: str | None) -> tuple[pd.DataFrame, list[str], list[str]]:
    if exposure is None:
        return pd.DataFrame(index=complete.index), [], []

    if exposure == "pulse_group":
        x = categorical_dummies(
            complete["pulse_rate_60_group"],
            PULSE_LABELS,
            PULSE_REFERENCE,
            "pulse_rate_60_group",
        )
        return x, ["pulse_rate_60_group"], ["pulse_rate_60_group"]

    if exposure == "pulse_continuous":
        return (
            complete[["pulse_rate_60_per_10bpm"]].copy(),
            ["pulse_rate_60_per_10bpm"],
            ["pulse_rate_60_per_10bpm"],
        )

    raise ValueError(f"Unknown exposure: {exposure}")


def adjusted_covariates(
    complete: pd.DataFrame,
    adjustment: str,
    include_sex: bool,
) -> tuple[pd.DataFrame, list[str], list[str]]:
    columns = []
    terms = []
    variables = []

    if adjustment in {"model1", "model2", "model3"}:
        columns.append(complete[["age_numeric"]])
        terms.append("age_numeric")
        variables.append("age_numeric")

        if include_sex:
            columns.append(categorical_dummies(complete["sex_group"], SEX_LEVELS, "Male", "sex"))
            terms.append("sex_group")
            variables.append("sex_group")

    if adjustment in {"model1", "model2", "model3"}:
        columns.extend([
            categorical_dummies(
                complete["education_group"],
                EDUCATION_LEVELS,
                "College or above",
                "education",
            ),
            categorical_dummies(complete["income_group"], INCOME_LEVELS, "High", "income"),
            categorical_dummies(complete["marital_status"], MARITAL_LEVELS, "Yes", "marital_status"),
            categorical_dummies(complete["working_status"], WORKING_LEVELS, "Yes", "working_status"),
        ])
        terms.extend([
            "education_group",
            "income_group",
            "marital_status",
            "working_status",
        ])
        variables.extend(terms[-4:])

    if adjustment in {"model2", "model3"}:
        columns.extend([
            categorical_dummies(complete["alcohol_group"], ALCOHOL_LEVELS, "Never", "alcohol"),
            categorical_dummies(complete["smoking_group"], SMOKING_LEVELS, "Never", "smoking"),
            categorical_dummies(
                complete["physical_activity_group"],
                PHYSICAL_ACTIVITY_LEVELS,
                ">=600 MET-min/week",
                "physical_activity",
            ),
            complete[["sleep_avg_weighted", "BMI_numeric"]],
        ])
        terms.extend([
            "alcohol_group",
            "smoking_group",
            "physical_activity_group",
            "sleep_avg_weighted",
            "BMI_numeric",
        ])
        variables.extend(terms[-5:])

    if not columns:
        return pd.DataFrame(index=complete.index), [], []

    return pd.concat(columns, axis=1), terms, variables


def model_design(
    data: pd.DataFrame,
    model_name: str,
    include_sex: bool,
) -> tuple[pd.Series, pd.DataFrame, str]:
    model_specs = {
        "model1_sociodemographic": ("model1", None),
        "model2_health_behavior": ("model2", None),
        "model3_bpm_group": ("model3", "pulse_group"),
    }

    if model_name not in model_specs:
        raise ValueError(f"Unknown model: {model_name}")

    adjustment, exposure = model_specs[model_name]
    exposure_columns = {
        None: [],
        "pulse_group": ["pulse_rate_60_group"],
        "pulse_continuous": ["pulse_rate_60_per_10bpm"],
    }[exposure]

    adjustment_columns = {
        "model1": [
            "age_numeric",
            "education_group",
            "income_group",
            "marital_status",
            "working_status",
        ],
        "model2": [
            "age_numeric",
            "education_group",
            "income_group",
            "marital_status",
            "working_status",
            "alcohol_group",
            "smoking_group",
            "physical_activity_group",
            "sleep_avg_weighted",
            "BMI_numeric",
        ],
        "model3": [
            "age_numeric",
            "education_group",
            "income_group",
            "marital_status",
            "working_status",
            "alcohol_group",
            "smoking_group",
            "physical_activity_group",
            "sleep_avg_weighted",
            "BMI_numeric",
        ],
    }[adjustment]
    if include_sex:
        adjustment_columns = adjustment_columns[:1] + ["sex_group"] + adjustment_columns[1:]

    variables = ["PHQ_binary"] + exposure_columns + adjustment_columns
    complete = data[variables].dropna()

    exposure_x, exposure_terms, _ = exposure_design(complete, exposure)
    adjustment_x, adjustment_terms, _ = adjusted_covariates(complete, adjustment, include_sex)
    x = pd.concat([exposure_x, adjustment_x], axis=1)
    description = "PHQ_binary ~ " + " + ".join(exposure_terms + adjustment_terms)

    y = complete["PHQ_binary"].astype(int)
    return y, add_constant(x), description


def article_style_model_design(
    data: pd.DataFrame,
    model_name: str,
) -> tuple[pd.Series, pd.DataFrame]:
    model_covariates = {
        "Model 1": ["age_numeric", "sex_group"],
        "Model 2": [
            "age_numeric",
            "sex_group",
            "education_group",
            "income_group",
            "alcohol_group",
            "smoking_group",
            "marital_status",
            "working_status",
        ],
        "Model 3": [
            "age_numeric",
            "sex_group",
            "education_group",
            "income_group",
            "alcohol_group",
            "smoking_group",
            "marital_status",
            "working_status",
            "BMI_numeric",
            "physical_activity_group",
        ],
    }

    variables = ["PHQ_binary", "pulse_rate_60_group"] + model_covariates[model_name]
    complete = data[variables].dropna()

    pulse_x, _, _ = exposure_design(complete, "pulse_group")
    columns = [
        pulse_x,
        complete[["age_numeric"]],
        categorical_dummies(complete["sex_group"], SEX_LEVELS, "Male", "sex"),
    ]

    if model_name in {"Model 2", "Model 3"}:
        columns.extend([
            categorical_dummies(
                complete["education_group"],
                EDUCATION_LEVELS,
                "College or above",
                "education",
            ),
            categorical_dummies(complete["income_group"], INCOME_LEVELS, "High", "income"),
            categorical_dummies(complete["alcohol_group"], ALCOHOL_LEVELS, "Never", "alcohol"),
            categorical_dummies(complete["smoking_group"], SMOKING_LEVELS, "Never", "smoking"),
            categorical_dummies(complete["marital_status"], MARITAL_LEVELS, "Yes", "marital_status"),
            categorical_dummies(complete["working_status"], WORKING_LEVELS, "Yes", "working_status"),
        ])

    if model_name == "Model 3":
        columns.extend([
            complete[["BMI_numeric"]],
            categorical_dummies(
                complete["physical_activity_group"],
                PHYSICAL_ACTIVITY_LEVELS,
                ">=600 MET-min/week",
                "physical_activity",
            ),
        ])

    y = complete["PHQ_binary"].astype(int)
    x = add_constant(pd.concat(columns, axis=1))
    return y, x


def fit_logit(y: pd.Series, x: pd.DataFrame):
    return sm.GLM(
        y,
        x,
        family=sm.families.Binomial(link=sm.families.links.Logit()),
    ).fit()


def extract_results(result, model_name: str, description: str, nobs: int) -> pd.DataFrame:
    params = result.params
    conf = result.conf_int()
    rows = []
    outcome = PHQ_CATEGORIES[1]
    reference_outcome = PHQ_CATEGORIES[0]

    for term in params.index:
        beta = params.loc[term]
        ci_low = conf.loc[term, 0]
        ci_high = conf.loc[term, 1]

        rows.append({
            "model": model_name,
            "formula": description,
            "n": int(nobs),
            "outcome_comparison": f"{outcome} vs {reference_outcome}",
            "term": term,
            "beta": beta,
            "std_error": result.bse.loc[term],
            "z": result.tvalues.loc[term],
            "p_value": result.pvalues.loc[term],
            "OR": np.exp(beta),
            "OR_95CI_low": np.exp(ci_low),
            "OR_95CI_high": np.exp(ci_high),
        })

    return pd.DataFrame(rows)


def format_p(value: float) -> str:
    if pd.isna(value):
        return ""
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def reference_rows_for_results(results: pd.DataFrame) -> pd.DataFrame:
    rows = []

    for _, model_row in results[
        ["analysis_group", "model", "formula", "n", "outcome_comparison"]
    ].drop_duplicates().iterrows():
        formula = model_row["formula"]

        for variable, (term_prefix, reference_level) in REFERENCE_ROWS.items():
            if variable not in formula:
                continue

            rows.append({
                "analysis_group": model_row["analysis_group"],
                "model": model_row["model"],
                "formula": formula,
                "n": int(model_row["n"]),
                "outcome_comparison": model_row["outcome_comparison"],
                "term": f"{term_prefix}[{reference_level}]",
                "beta": 0.0,
                "std_error": np.nan,
                "z": np.nan,
                "p_value": np.nan,
                "OR": 1.0,
                "OR_95CI_low": np.nan,
                "OR_95CI_high": np.nan,
                "is_reference": True,
            })

    return pd.DataFrame(rows)


def add_reference_rows(results: pd.DataFrame) -> pd.DataFrame:
    table = results.copy()
    table["is_reference"] = False
    reference_rows = reference_rows_for_results(table)

    if reference_rows.empty:
        return table

    return sort_result_rows(pd.concat([table, reference_rows], ignore_index=True))


def term_sort_key(term: str) -> tuple[int, int, str]:
    if term in TERM_ORDER:
        order, level_order = TERM_ORDER[term]
        return order, level_order, term

    if "[" in term and term.endswith("]"):
        prefix, level = term[:-1].split("[", 1)
        base_order, _ = TERM_ORDER.get(prefix, (99, 0))
        levels = LEVEL_ORDER.get(prefix, [])
        level_order = levels.index(level) if level in levels else 99
        return base_order, level_order, term

    return 99, 0, term


def sort_result_rows(results: pd.DataFrame) -> pd.DataFrame:
    table = results.copy()
    sort_keys = table["term"].map(term_sort_key)
    table["_term_order"] = sort_keys.map(lambda value: value[0])
    table["_level_order"] = sort_keys.map(lambda value: value[1])
    table["_term_label"] = sort_keys.map(lambda value: value[2])
    table = table.sort_values(
        ["analysis_group", "model", "_term_order", "_level_order", "_term_label"],
        kind="stable",
    )
    return table.drop(columns=["_term_order", "_level_order", "_term_label"])


def make_readable_table(results: pd.DataFrame) -> pd.DataFrame:
    table = results.copy()
    table["OR (95% CI)"] = table.apply(
        lambda row: (
            "1.00 (Reference)"
            if row.get("is_reference", False)
            else (
                f"{row['OR']:.2f} "
                f"({row['OR_95CI_low']:.2f}-{row['OR_95CI_high']:.2f})"
            )
        ),
        axis=1,
    )
    table["p-value"] = table["p_value"].map(format_p)
    return table[
        [
            "analysis_group",
            "model",
            "n",
            "outcome_comparison",
            "term",
            "OR (95% CI)",
            "p-value",
        ]
    ]


def make_bpm_summary_table(main_results: pd.DataFrame) -> pd.DataFrame:
    bpm_results = main_results[main_results["model"] == "model3_bpm_group"].copy()
    rows = []

    for label in PULSE_LABELS:
        row = {
            "BPM group": label,
            "Reference": PULSE_REFERENCE if label == PULSE_REFERENCE else "",
        }

        for analysis_group in ["Total"]:
            if label == PULSE_REFERENCE:
                row[f"{analysis_group} OR (95% CI)"] = "1.00 (Reference)"
                row[f"{analysis_group} p-value"] = ""
                continue

            term = f"pulse_rate_60_group[{label}]"
            match = bpm_results[
                (bpm_results["analysis_group"] == analysis_group)
                & (bpm_results["term"] == term)
            ]

            if match.empty:
                row[f"{analysis_group} OR (95% CI)"] = ""
                row[f"{analysis_group} p-value"] = ""
                continue

            result_row = match.iloc[0]
            row[f"{analysis_group} OR (95% CI)"] = (
                f"{result_row['OR']:.2f} "
                f"({result_row['OR_95CI_low']:.2f}-{result_row['OR_95CI_high']:.2f})"
            )
            row[f"{analysis_group} p-value"] = format_p(result_row["p_value"])

        rows.append(row)

    return pd.DataFrame(rows)


def make_article_style_rhr_table(data: pd.DataFrame) -> tuple[pd.DataFrame, dict[tuple[str, str], bool]]:
    article_models = ["Model 1", "Model 2", "Model 3"]
    model_results = {}
    significant_cells = {}

    analysis_groups = {"Total": data.copy()}

    for analysis_label, group_data in analysis_groups.items():
        for model_name in article_models:
            y, x = article_style_model_design(group_data, model_name)
            model_results[(analysis_label, model_name)] = fit_logit(y, x)

    rows = []
    for analysis_label, group_data in analysis_groups.items():
        count_data = group_data[["PHQ_binary", "pulse_rate_60_group"]].dropna()

        for pulse_label in PULSE_LABELS:
            pulse_data = count_data[count_data["pulse_rate_60_group"] == pulse_label]
            case_count = int(pulse_data["PHQ_binary"].sum())
            total_count = int(len(pulse_data))
            row = {
                "Analysis group": analysis_label,
                "RHR (bpm)": pulse_label,
                "Number of case": f"{case_count}/{total_count}",
            }

            for model_name in article_models:
                if pulse_label == PULSE_REFERENCE:
                    row[model_name] = "1"
                    significant_cells[(analysis_label, pulse_label, model_name)] = False
                    continue

                term = f"pulse_rate_60_group[{pulse_label}]"
                result = model_results[(analysis_label, model_name)]
                beta = result.params.loc[term]
                ci_low, ci_high = result.conf_int().loc[term]
                p_value = result.pvalues.loc[term]

                row[model_name] = (
                    f"{np.exp(beta):.2f} "
                    f"({np.exp(ci_low):.2f}-{np.exp(ci_high):.2f})"
                )
                significant_cells[(analysis_label, pulse_label, model_name)] = p_value < 0.05

            rows.append(row)

    return pd.DataFrame(rows), significant_cells


def write_model_123_excel_file(results: pd.DataFrame) -> Path:
    from openpyxl.styles import Alignment, Font, PatternFill

    model_sheet_names = {
        "model1_sociodemographic": "Model 1",
        "model2_health_behavior": "Model 2",
        "model3_bpm_group": "Model 3",
    }
    font = Font(name="Times New Roman", size=14)
    header_font = Font(name="Times New Roman", size=14, bold=True)
    alignment = Alignment(vertical="center")
    header_fill = PatternFill("solid", fgColor="D9EAD3")

    with pd.ExcelWriter(MODEL_123_XLSX, engine="openpyxl") as writer:
        for model_name, sheet_name in model_sheet_names.items():
            model_results = results[results["model"] == model_name].copy()
            make_readable_table(model_results).to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"

            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = alignment

            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = font
                    cell.alignment = alignment

            for column_cells in worksheet.columns:
                max_length = max(
                    len(str(cell.value)) if cell.value is not None else 0
                    for cell in column_cells
                )
                worksheet.column_dimensions[column_cells[0].column_letter].width = min(
                    max(max_length + 2, 12),
                    60,
                )

    return MODEL_123_XLSX


def keep_main_exposure_terms(results: pd.DataFrame) -> pd.DataFrame:
    return results[
        results["term"].str.startswith("pulse_rate_60_group[")
        | (results["term"] == "pulse_rate_60_per_10bpm")
    ].copy()


def clean_old_outputs() -> None:
    if not OUT_DIR.exists():
        return

    for path in OUT_DIR.iterdir():
        if path.is_dir():
            shutil.rmtree(path)
        else:
            path.unlink()


def phq_binary_verification_table(data: pd.DataFrame) -> pd.DataFrame:
    return pd.crosstab(
        data["PHQ_group"],
        data["PHQ_binary"],
        dropna=False,
    )


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    clean_old_outputs()
    data = prepare_data(DATA_PATH)

    model_names = [
        "model1_sociodemographic",
        "model2_health_behavior",
        "model3_bpm_group",
    ]

    analysis_groups = {"Total": data.copy()}

    all_results = []

    for analysis_group, group_data in analysis_groups.items():
        for model_name in model_names:
            y, x, description = model_design(group_data, model_name, include_sex=True)
            result = fit_logit(y, x)

            result_table = extract_results(result, model_name, description, len(y))
            result_table.insert(0, "analysis_group", analysis_group)

            all_results.append(result_table)

    combined = pd.concat(all_results, ignore_index=True)
    combined_with_references = add_reference_rows(combined)
    output_path = write_model_123_excel_file(combined_with_references)

    print('Verification: table(PHQ_group, PHQ_binary, useNA = "ifany")')
    print(phq_binary_verification_table(data))
    print("Saved regression output to:")
    print(output_path)


if __name__ == "__main__":
    main()
