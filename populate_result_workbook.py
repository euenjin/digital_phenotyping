from pathlib import Path
import shutil

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


RESULT_PATH = Path("result.xlsx")
BACKUP_PATH = Path("result_before_refill.xlsx")
DATA_PATH = Path("df_cohort_cleaned.csv")
TABLE1_PATH = Path("table1_characteristics_by_heart_rate.xlsx")
REGRESSION_PATH = Path("regression_results/model_1_2_3.xlsx")

PULSE_LABELS = ["<60 bpm", "60-69 bpm", "70-79 bpm", "80-89 bpm", ">=90 bpm"]
PULSE_REFERENCE = "<60 bpm"


def format_p(value):
    if pd.isna(value) or value == "":
        return ""
    if isinstance(value, str):
        return value
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def prepare_data() -> pd.DataFrame:
    df = pd.read_csv(DATA_PATH)
    df["PHQ_binary"] = df["PHQ_group"].map({"Minimal": 0, "Mild_or_greater": 1})
    df["pulse_rate_60_group"] = pd.cut(
        df["pulse_rate_60"],
        bins=[-float("inf"), 59, 69, 79, 89, float("inf")],
        labels=PULSE_LABELS,
        right=True,
    )
    return df


def load_regression_results() -> pd.DataFrame:
    model_labels = {"Model 1": "Model 1", "Model 2": "Model 2", "Model 3": "Model 3"}
    frames = []
    for sheet_name, model_label in model_labels.items():
        frame = pd.read_excel(REGRESSION_PATH, sheet_name=sheet_name)
        frame.insert(0, "Model label", model_label)
        frames.append(frame)
    return pd.concat(frames, ignore_index=True)


def make_criteria_table(df: pd.DataFrame) -> pd.DataFrame:
    years = sorted(df["year"].dropna().astype(int).unique())
    return pd.DataFrame(
        [
            ("Study topic", "Resting heart rate and PHQ-9 depressive symptoms among Korean adults"),
            ("Data source", f"KNHANES even survey years {years[0]}-{years[-1]} ({', '.join(map(str, years))})"),
            ("Final analytic cohort", f"{len(df):,} participants"),
            ("Outcome", "Mild or greater depressive symptoms vs minimal symptoms"),
            ("Outcome coding", "PHQ-9 0-4 = Minimal; PHQ-9 >=5 = Mild_or_greater"),
            ("Main exposure", "Resting heart rate measured for 60 seconds"),
            ("Exposure categories", "<60, 60-69, 70-79, 80-89, >=90 bpm"),
            ("Reference exposure", "<60 bpm"),
            ("Primary analysis", "Logistic regression; odds ratio (OR) and 95% confidence interval"),
            ("Model 1", "Sociodemographic variables: age, sex, education, household income, marital status, working status"),
            ("Model 2", "Model 1 + health behavior variables: alcohol, smoking, physical activity, sleep duration, BMI"),
            ("Model 3", "Model 2 + resting heart rate group"),
            ("Sex handling", "Sex is included as a demographic covariate; odds ratios are not stratified by sex"),
            ("Exclusion summary", "See the preserved 연구대상자 sheet for cohort flow"),
        ],
        columns=["Item", "Definition"],
    )


def make_model_summary_table(regression: pd.DataFrame) -> pd.DataFrame:
    n_by_model = regression.groupby("Model label")["n"].first().to_dict()
    return pd.DataFrame(
        [
            (
                "Model 1",
                "Sociodemographic",
                "age, sex, education, income, marital status, working status",
                f"{int(n_by_model['Model 1']):,}",
            ),
            (
                "Model 2",
                "Sociodemographic + health behavior",
                "Model 1 + alcohol, smoking, physical activity, sleep duration, BMI",
                f"{int(n_by_model['Model 2']):,}",
            ),
            (
                "Model 3",
                "Sociodemographic + health behavior + RHR group",
                "Model 2 + resting heart rate group",
                f"{int(n_by_model['Model 3']):,}",
            ),
        ],
        columns=["Model", "Label", "Adjusted covariates", "Complete-case N"],
    )


def make_rhr_or_summary(df: pd.DataFrame, regression: pd.DataFrame) -> pd.DataFrame:
    count_data = df[["PHQ_binary", "pulse_rate_60_group"]].dropna()
    rows = []
    for pulse_label in PULSE_LABELS:
        pulse_data = count_data[count_data["pulse_rate_60_group"] == pulse_label]
        row = {
            "RHR group": pulse_label,
            "Cases / total": f"{int(pulse_data['PHQ_binary'].sum()):,}/{len(pulse_data):,}",
        }
        if pulse_label == PULSE_REFERENCE:
            row["Model 3 OR (95% CI)"] = "1.00 (Reference)"
            row["Model 3 p-value"] = ""
            rows.append(row)
            continue

        term = f"pulse_rate_60_group[{pulse_label}]"
        match = regression[
            (regression["Model label"] == "Model 3")
            & (regression["term"] == term)
        ]
        if match.empty:
            row["Model 3 OR (95% CI)"] = ""
            row["Model 3 p-value"] = ""
            rows.append(row)
            continue

        result = match.iloc[0]
        row["Model 3 OR (95% CI)"] = result["OR (95% CI)"]
        row["Model 3 p-value"] = format_p(result["p-value"])
        rows.append(row)
    return pd.DataFrame(rows)


def make_phq_cross_table(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    total_n = len(df)
    total_cases = int(df["PHQ_binary"].sum())
    rows.append(
        {
            "RHR group": "Total",
            "Total N": f"{total_n:,}",
            "Minimal, n (%)": f"{total_n - total_cases:,} ({(total_n - total_cases) / total_n * 100:.1f})",
            "Mild_or_greater, n (%)": f"{total_cases:,} ({total_cases / total_n * 100:.1f})",
        }
    )
    for pulse_label in PULSE_LABELS:
        pulse_data = df[df["pulse_rate_60_group"] == pulse_label]
        n = len(pulse_data)
        cases = int(pulse_data["PHQ_binary"].sum())
        rows.append(
            {
                "RHR group": pulse_label,
                "Total N": f"{n:,}",
                "Minimal, n (%)": f"{n - cases:,} ({(n - cases) / n * 100:.1f})",
                "Mild_or_greater, n (%)": f"{cases:,} ({cases / n * 100:.1f})",
            }
        )
    return pd.DataFrame(rows)


def make_variable_dictionary() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ("PHQ_group", "Depressive symptom group", "Minimal / Mild_or_greater"),
            ("PHQ_binary", "Regression outcome", "0 = Minimal, 1 = Mild_or_greater"),
            ("pulse_rate_60", "Resting heart rate", "beats per minute"),
            ("pulse_rate_60_group", "Categorical resting heart rate", "<60 bpm reference"),
            ("age_numeric", "Age", "years"),
            ("sex_numeric / sex_group", "Sex", "1 = Male, 2 = Female; Male reference"),
            ("household_income / income_group", "Household income quartile", "High reference"),
            ("edu_numeric / education_group", "Education level", "College or above reference"),
            ("marri_1_numeric / marital_status", "Marital status", "Yes reference"),
            ("employment_status / working_status", "Working status", "Yes reference"),
            ("smoking_history / smoking_group", "Smoking history", "Never reference"),
            ("BD1_11_numeric / alcohol_group", "Alcohol use", "Never reference"),
            ("physical_activity_group", "Total physical activity", ">=600 MET-min/week reference"),
            ("sleep_avg_weighted", "Average sleep duration", "Weighted weekday/weekend hours"),
            ("BMI_numeric", "Body mass index", "kg/m2"),
        ],
        columns=["Variable", "Description", "Coding / note"],
    )


def reset_sheet(wb: openpyxl.Workbook, title: str):
    if title in wb.sheetnames:
        del wb[title]
    return wb.create_sheet(title)


def write_dataframe(ws, df: pd.DataFrame, start_row: int = 1, start_col: int = 1):
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start_row):
        for c_idx, value in enumerate(row, start_col):
            ws.cell(r_idx, c_idx, value)


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    subheader_fill = PatternFill("solid", fgColor="F3F6F1")
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(bottom=thin_gray)

    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Malgun Gothic", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
        if row and row[0].row == 1:
            for cell in row:
                cell.font = Font(name="Malgun Gothic", size=10, bold=True)
                cell.fill = header_fill
        elif row and row[0].value and all(cell.value in (None, "") for cell in row[1:]):
            for cell in row:
                cell.fill = subheader_fill
                cell.font = Font(name="Malgun Gothic", size=10, bold=True)

    for column_cells in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 55)


def main() -> None:
    if not BACKUP_PATH.exists():
        shutil.copy2(RESULT_PATH, BACKUP_PATH)

    df = prepare_data()
    regression = load_regression_results()
    table1 = pd.read_excel(TABLE1_PATH, sheet_name="Total")

    wb = openpyxl.load_workbook(RESULT_PATH)
    for sheet_name in list(wb.sheetnames):
        if sheet_name != "연구대상자":
            del wb[sheet_name]

    sheet_data = [
        ("기준", make_criteria_table(df)),
        ("모델요약", make_model_summary_table(regression)),
        ("T1_특성_RHR", table1),
        ("T2_심박수_OR", make_rhr_or_summary(df, regression)),
        ("T3_전체모델_OR", regression),
        ("T4_PHQ_RHR_교차표", make_phq_cross_table(df)),
        ("변수정의", make_variable_dictionary()),
    ]

    for title, frame in sheet_data:
        ws = reset_sheet(wb, title)
        write_dataframe(ws, frame)
        style_sheet(ws)

    wb.active = wb.sheetnames.index("연구대상자")
    wb.save(RESULT_PATH)


if __name__ == "__main__":
    main()
